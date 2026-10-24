"""
Extract all text slides and images from the GN Construction PPT.
"""
from pptx import Presentation
from pptx.util import Inches
import os
import json
from PIL import Image
import io

PPTX_PATH = r"C:\Users\nwabl\.gemini\antigravity\brain\aa9ed5fb-6156-453e-a1a5-011eea8555c6\.user_uploaded\media_1792765634112_a150fed4.pptx"
OUT_DIR = r"C:\Users\nwabl\Downloads\Gn Construction website\gn_ppt_assets"
IMAGES_DIR = os.path.join(OUT_DIR, "images")
os.makedirs(IMAGES_DIR, exist_ok=True)

prs = Presentation(PPTX_PATH)

slides_data = []
img_count = 0

for slide_num, slide in enumerate(prs.slides, start=1):
    slide_texts = []
    slide_images = []
    
    for shape in slide.shapes:
        # Extract text
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                text = para.text.strip()
                if text:
                    slide_texts.append(text)
        
        # Extract images
        if shape.shape_type == 13:  # MSO_SHAPE_TYPE.PICTURE
            img_count += 1
            image = shape.image
            ext = image.ext
            fname = f"slide{slide_num:02d}_img{img_count:02d}.{ext}"
            fpath = os.path.join(IMAGES_DIR, fname)
            with open(fpath, "wb") as f:
                f.write(image.blob)
            
            # Get dimensions
            try:
                pil_img = Image.open(io.BytesIO(image.blob))
                w, h = pil_img.size
            except:
                w, h = 0, 0
            
            slide_images.append({
                "file": fname,
                "width": w,
                "height": h,
                "size_bytes": len(image.blob)
            })
            print(f"  Slide {slide_num}: Image -> {fname} ({w}x{h})")
    
    slides_data.append({
        "slide": slide_num,
        "texts": slide_texts,
        "images": slide_images
    })
    
    if slide_texts:
        print(f"Slide {slide_num}: {slide_texts[:3]}")

# Save all text data
with open(os.path.join(OUT_DIR, "slides_data.json"), "w", encoding="utf-8") as f:
    json.dump(slides_data, f, indent=2, ensure_ascii=False)

print(f"\n=== DONE ===")
print(f"Total slides: {len(slides_data)}")
print(f"Total images extracted: {img_count}")
print(f"Output dir: {OUT_DIR}")
